"""Final verification of the summary document, written to a UTF-8 file so the
results can be read back without the console mangling Chinese text."""

import json
from docx import Document
from docx.oxml.ns import qn

DOC = "蓝鲸小深-项目总结.docx"
OUT = "build/doc-verify.txt"

doc = Document(DOC)

lines = []
lines.append("=== 中文完整性 ===")
full = "\n".join(p.text for p in doc.paragraphs)
for t in doc.tables:
    for row in t.rows:
        for c in row.cells:
            full += "\n" + c.text

for probe in ["蓝鲸小深", "边缘固定", "隐私边界", "逆向设计", "命中测试",
              "GooglePiggy", "Codex Pets", "判断不成立", "GetAsyncKeyState"]:
    lines.append(f"  {probe:<20} {'存在' if probe in full else '缺失'}")
lines.append(f"  总字符数: {len(full)}")
lines.append(f"  可疑乱码: {sum(full.count(c) for c in '锛鈥鉁鏂鐩锟')}")

lines.append("")
lines.append("=== 字体（eastAsia 必须显式设置，否则中文回退）===")
missing = 0
total_runs = 0
for p in doc.paragraphs:
    for r in p.runs:
        if not r.text.strip():
            continue
        total_runs += 1
        rpr = r._element.find(qn("w:rPr"))
        ea = None
        if rpr is not None:
            rf = rpr.find(qn("w:rFonts"))
            if rf is not None:
                ea = rf.get(qn("w:eastAsia"))
        if ea is None:
            missing += 1
lines.append(f"  非空 run 总数: {total_runs}")
lines.append(f"  未设置 eastAsia 的 run: {missing}")

lines.append("")
lines.append("=== 标题结构 ===")
for p in doc.paragraphs:
    if p.style.name.startswith("Heading"):
        lvl = p.style.name.replace("Heading ", "")
        indent = "  " * (int(lvl) - 1 if lvl.isdigit() else 0)
        lines.append(f"  {indent}{p.text}")

lines.append("")
lines.append("=== 表格 ===")
for i, t in enumerate(doc.tables, 1):
    hdr = " | ".join(c.text for c in t.rows[0].cells)
    lines.append(f"  表{i}: {len(t.rows)}行 x {len(t.columns)}列  表头: {hdr}")

lines.append("")
lines.append("=== 段落/字数统计 ===")
lines.append(f"  段落数: {len(doc.paragraphs)}")
lines.append(f"  表格数: {len(doc.tables)}")

with open(OUT, "w", encoding="utf-8") as fh:
    fh.write("\n".join(lines))
print(f"wrote {OUT}")
